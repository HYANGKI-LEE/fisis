"""card_data.py(FISIS 계산값)를 한기평 신용카드 재무자료(엑셀) 값과 대조한다.

사용: python scripts/verify_cards.py [엑셀경로] [기간 예: 202606]
"""
import sys
import warnings
from pathlib import Path

import numpy as np
import openpyxl

warnings.filterwarnings("ignore")
sys.path.insert(0, str(Path(__file__).resolve().parent))
from card_data import CardData  # noqa: E402

XL_NAMES = {"신한카드": "신한카드", "삼성카드": "삼성카드", "케이비국민카드": "KB국민카드", "현대카드": "현대카드", "롯데카드": "롯데카드", "우리카드": "우리카드", "하나카드": "하나카드"}
BS = {6: "bs_cashdep", 7: "bs_sec", 8: "bs_ops", 9: "bs_card", 10: "bs_inst", 11: "bs_lease", 12: "bs_loan", 13: "bs_nbt", 14: "bs_oth_a",
      15: "bs_asset", 16: "bs_borrow", 17: "bs_b_loan", 18: "bs_bond", 20: "bs_pay", 21: "bs_unused", 22: "bs_liab", 23: "bs_capital", 24: "bs_equity"}
IS = {27: "rev", 28: "r_card", 29: "r_inst", 30: "r_lease", 31: "r_loanint", 32: "r_nbt", 33: "r_sec", 35: "cost", 36: "c_int", 37: "c_card",
      38: "c_inst", 39: "c_lease", 40: "c_nbt", 41: "c_sga", 42: "c_sec", 44: "ppop", 45: "dae", 48: "op", 49: "ni"}


def main():
    xlsx = Path(sys.argv[1]) if len(sys.argv) > 1 and sys.argv[1].endswith(".xlsx") else Path(r"C:\Users\infomax\Downloads\(Data Sheet) 신용카드 재무자료_2026.2Q.xlsx")
    ym = sys.argv[2] if len(sys.argv) > 2 else "202606"
    ws = openpyxl.load_workbook(xlsx, data_only=True)[f"{ym[:4]}({ym[4:]})"]
    cols = {}
    for c in range(6, ws.max_column + 1):
        v = ws.cell(26, c).value
        if isinstance(v, str) and v.strip() in XL_NAMES:
            cols[XL_NAMES[v.strip()]] = c
    cd = CardData()
    bad = 0
    for r, key in list(BS.items()) + [(r, k) for r, k in IS.items()]:
        is_bs = r in BS
        diffs = []
        for co, c in cols.items():
            ex = ws.cell(r, c).value
            if not isinstance(ex, (int, float)):
                continue
            mine = cd.metrics[f"bs_{key[3:]}"]["df"][ym].get(co) if is_bs else cd.IS_y[key][ym].get(co)
            if is_bs:
                mine = cd.metrics[key]["df"][ym].get(co)
            if mine is None or not np.isfinite(mine) or abs(mine - ex) > max(1.0, abs(ex) * 0.0):
                diffs.append((co, ex, None if mine is None or not np.isfinite(mine) else round(float(mine), 1)))
        bad += bool(diffs)
        print(f"{'OK ' if not diffs else 'BAD'} r{r:<3d} {str(ws.cell(r, 5).value).strip()[:16]:18s} {key:10s} {diffs[:4] if diffs else ''}")
    print("불일치 행:", bad)


if __name__ == "__main__":
    main()
